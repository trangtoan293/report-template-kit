"""Xuất SQL Oracle, file Excel đặc tả, và điền thử template bằng engine rt_office.py."""
import json
import subprocess
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from customsso import rt_office  # python3 sync_engine.py . <values-report-builder.yaml>
from make_data import AMT6, OUT, REPORTS, SAMPLES, dataset_columns, sample_row, source_tables

DATE_COLS = {"tu_ngay", "den_ngay", "ngay_so_lieu"}


# ================================================================ SQL
def agg(expr, src, where, order):
    # Nối nhiều dòng thành 1 CLOB, phân cách CHR(10). LISTAGG bị giới hạn 4000 byte nên dùng XMLAGG.
    return (f"(SELECT RTRIM(DBMS_XMLGEN.CONVERT(XMLAGG(XMLELEMENT(e, {expr} || CHR(10)) ORDER BY {order})"
            f".EXTRACT('//text()').getClobVal(), 1), CHR(10))\n       FROM {src} c WHERE {where})")


def hdr_keys(r):
    return [col for tbl, col, _, pk, _ in source_tables(r) if tbl.endswith("_HDR") and pk == "PK"]


def join(r, alias="c"):
    return " AND ".join(f"{alias}.{k} = h.{k}" for k in hdr_keys(r))


def expr_for(r, name):
    code, ct, w = r["code"], f"RPT_{r['code']}_CT", join(r)
    dec = "CASE WHEN h.loai_tien = 'VND' THEN 0 ELSE 2 END"
    common = {"ngay_ky": "TO_CHAR(h.ngay_lap, 'DD')", "thang_ky": "TO_CHAR(h.ngay_lap, 'MM')",
              "nam_ky": "TO_CHAR(h.ngay_lap, 'YYYY')", "thang_bc": "h.ky_bc",
              "ngay_tao": "TO_CHAR(h.ngay_tao, 'DD/MM/YYYY HH24:MI')"}
    if name in common:
        return common[name]
    if name in DATE_COLS:
        return f"TO_CHAR(h.{name}, 'DD/MM/YYYY')"
    if code == "BM07" and name in ("du_dau_ky", "tong_thu", "tong_chi", "du_cuoi_ky"):
        return f"F_RPT_SO(h.{name}, {dec})"
    if code == "BM04":
        if name.startswith("tong_"):
            return f"F_RPT_SO(h.{name})"
        if name[:2] in ("a_", "b_"):
            p, f = name[0], name[2:]
            mx = {"ten_tk": 48, "tk_nhnn": 10, "tk_fcc": 12}
            e = f"SUBSTR(c.{f}, 1, {mx[f]})" if f in mx else f"F_RPT_SO(c.{f})"
            return agg(e, ct, f"{w} AND c.phan = '{p.upper()}'", "c.stt")
    if code == "BM05" and name[:3] in ("ct_", "tc_"):
        f = name[3:]
        if name.startswith("ct_"):
            return agg("c.nhan" if f == "loai_tien" else f"F_RPT_SO(c.{f})", "V_RPT_BM05_DONG", w, "c.ma_cn, c.thu_tu, c.nhan")
        return agg("'Tổng cộng ' || c.loai_tien" if f == "loai_tien" else f"F_RPT_SO(c.{f})", "V_RPT_BM05_TONG", w, "c.loai_tien")
    if code == "BM06":
        if name.startswith("ct_"):
            f = name[3:]
            e = {"so_tk": "c.so_tk_ct", "ten_tk": "SUBSTR(c.ten_tk_ct, 1, 52)"}.get(f, f"F_RPT_SO(c.{f})")
            return agg(e, ct, w, "c.so_tk_ct")
        if name.startswith("tong_"):
            return f"(SELECT F_RPT_SO(SUM(c.{name[5:]})) FROM {ct} c WHERE {w})"
    if code == "BM07" and name.startswith("ct_"):
        f = name[3:]
        mx = {"so_ct": 20, "dien_giai": 50, "giao_dich_vien": 32, "so_ct_gdv": 20}
        e = f"SUBSTR(c.{f}, 1, {mx[f]})" if f in mx else f"F_RPT_SO(c.{f}, {dec})"
        return agg(e, ct, w, "c.stt")
    return f"h.{name}"


def ddl(r):
    out = []
    tables = OrderedDict()
    for tbl, col, typ, pk, desc in source_tables(r):
        tables.setdefault(tbl, []).append((col, typ, pk, desc))
    for tbl, cols in tables.items():
        # dấu phẩy đặt TRƯỚC comment "--" (mọi cột đều có phẩy vì dòng cuối là CONSTRAINT)
        body = "\n".join(f"  {c:<18} {t + (' NOT NULL' if pk else '') + ',':<24}{'-- ' + d if d else ''}".rstrip() for c, t, pk, d in cols)
        pks = ", ".join(c for c, _, pk, _ in cols if pk)
        out.append(f"CREATE TABLE {tbl} (\n{body}\n  CONSTRAINT PK_{tbl} PRIMARY KEY ({pks})\n);")
    return "\n\n".join(out)


def view_fixed(r):
    keys = hdr_keys(r)
    pivots = []
    for i in range(len(r["rows"])):
        for suf, src, _ in r["measures"]:
            name = f"{r['prefix']}_r{i + 1:02d}" + (f"_{suf}" if suf else "")
            pivots.append((name, f"MAX(CASE WHEN ma_chi_tieu = '{r['prefix']}_r{i + 1:02d}' THEN F_RPT_SO({src}, 0, '{r['neg']}') END)"))
    piv_names = {n for n, _ in pivots}
    sel = [f"  {expr_for(r, n) if n not in piv_names else 'p.' + n} AS {n}" for n, *_ in dataset_columns(r)]
    inner = ",\n".join(f"    {e} AS {n}" for n, e in pivots)
    return (f"CREATE OR REPLACE VIEW V_RPT_{r['code']} AS\nSELECT\n" + ",\n".join(sel) +
            f"\nFROM RPT_{r['code']}_HDR h\nLEFT JOIN (\n  SELECT {', '.join(keys)},\n{inner}\n  FROM RPT_{r['code']}_CT\n"
            f"  GROUP BY {', '.join(keys)}\n) p ON {join(r, 'p')};")


def view_list(r):
    sel = [f"  {expr_for(r, n)} AS {n}" for n, *_ in dataset_columns(r)]
    pre = ""
    if r["code"] == "BM05":
        k = "ma_tk, tu_ngay, den_ngay"
        nulls = ", ".join(f"CAST(NULL AS NUMBER) AS {f}" for f in AMT6)
        pre = (f"-- Dòng in: tiêu đề chi nhánh (thu_tu = 0) + các dòng loại tiền (thu_tu = 1)\n"
               f"CREATE OR REPLACE VIEW V_RPT_BM05_DONG AS\n"
               f"SELECT {k}, ma_cn, 0 AS thu_tu, SUBSTR(ma_cn || ' - ' || ten_cn, 1, 40) AS nhan, {nulls}\n"
               f"FROM RPT_BM05_CT GROUP BY {k}, ma_cn, ten_cn\nUNION ALL\n"
               f"SELECT {k}, ma_cn, 1, loai_tien, {', '.join(AMT6)} FROM RPT_BM05_CT;\n\n"
               f"CREATE OR REPLACE VIEW V_RPT_BM05_TONG AS\n"
               f"SELECT {k}, loai_tien, {', '.join(f'SUM({f}) AS {f}' for f in AMT6)}\n"
               f"FROM RPT_BM05_CT GROUP BY {k}, loai_tien;\n\n")
    return pre + f"CREATE OR REPLACE VIEW V_RPT_{r['code']} AS\nSELECT\n" + ",\n".join(sel) + f"\nFROM RPT_{r['code']}_HDR h;"


SQL_HEAD = """-- =====================================================================
-- CẤU TRÚC DỮ LIỆU CHO TEMPLATE REPORT BUILDER (Superset add-on)
-- Oracle 19c+. Mỗi báo cáo: bảng RPT_BMxx_HDR (1 dòng / 1 file in),
-- RPT_BMxx_CT (dòng chi tiết), view V_RPT_BMxx = dataset khai báo trên Superset.
-- Tên cột của view = tên biến {#...#} trong template -> Report Builder auto-map.
-- PostgreSQL: thay khối XMLAGG bằng string_agg(<expr>, E'\\n' ORDER BY <order>)
--             và F_RPT_SO bằng to_char(...) tương ứng.
-- Lưu ý: cột nhiều dòng là CLOB; nếu driver Superset trả CLOB không đúng dạng chuỗi,
--        bật MAX_STRING_SIZE=EXTENDED và dùng LISTAGG ... (tối đa 32767 byte).
-- =====================================================================

-- Định dạng số kiểu Việt Nam: 1234567.5 -> '1.234.567,50'; âm -> '-x' hoặc '(x)'; NULL -> NULL
CREATE OR REPLACE FUNCTION F_RPT_SO(p_so NUMBER, p_so_le NUMBER DEFAULT 0, p_am VARCHAR2 DEFAULT '-')
RETURN VARCHAR2 DETERMINISTIC IS
  v VARCHAR2(60);
BEGIN
  IF p_so IS NULL THEN RETURN NULL; END IF;
  v := TO_CHAR(ABS(p_so), 'FM999G999G999G999G999G990' || CASE WHEN p_so_le > 0 THEN 'D' || RPAD('0', p_so_le, '0') END,
               'NLS_NUMERIC_CHARACTERS='',.''');
  IF p_so < 0 THEN RETURN CASE p_am WHEN '()' THEN '(' || v || ')' ELSE '-' || v END; END IF;
  RETURN v;
END;
/
"""


def write_sql():
    parts = [SQL_HEAD]
    for r in REPORTS:
        parts.append(f"\n-- =====================================================================\n-- {r['code']}. {r['name']}\n"
                     f"-- Template: {r['file']}\n-- Dataset : V_RPT_{r['code']}  ({r['grain']})\n"
                     f"-- =====================================================================\n")
        parts.append(ddl(r))
        parts.append("\n\n" + (view_fixed(r) if r["kind"] == "fixed" else view_list(r)) + "\n")
    (OUT / "Cau_truc_du_lieu_Oracle.sql").write_text("\n".join(parts), encoding="utf-8")


# ================================================================ Excel
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def put_table(ws, top, headers, rows, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(top, j, h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, top + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.border, c.alignment = BOX, Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    return top + len(rows) + 2


def preview(v):
    s = "" if v is None else str(v)
    lines = s.split("\n")
    return s if len(lines) == 1 else " ⏎ ".join(lines[:3]) + f"  … ({len(lines)} dòng)"


GUIDE = [
    "QUY ƯỚC DỮ LIỆU CHO TEMPLATE REPORT BUILDER",
    "",
    "1. Biến trong template có dạng {#ten_bien#}. Report Builder tự ánh xạ biến với cột dataset CÙNG TÊN (không phân biệt hoa thường).",
    "2. 1 dòng dataset = 1 file in (PDF). Chọn nhiều dòng -> tải ZIP nhiều PDF. Engine KHÔNG hỗ trợ vòng lặp.",
    "3. Báo cáo cố định (BM01-BM03): mỗi ô chỉ tiêu là 1 cột -> dataset dạng 'ngang' (pivot từ bảng chi tiết, xem view SQL).",
    "4. Báo cáo danh sách (BM04-BM07): mỗi cột của bảng là 1 biến NHIỀU DÒNG - các giá trị nối bằng ký tự xuống dòng (CHR(10)).",
    "   - Tất cả cột nhiều dòng của cùng 1 bảng PHẢI có cùng số dòng và cùng thứ tự (giá trị rỗng vẫn giữ 1 dòng trống).",
    "   - Cột chữ phải cắt theo độ dài tối đa ghi trong đặc tả để không bị xuống dòng tự động (lệch hàng giữa các cột).",
    "5. Engine in giá trị nguyên văn (str) -> SỐ PHẢI ĐƯỢC ĐỊNH DẠNG SẴN thành chuỗi: dấu chấm phân cách nghìn, dấu phẩy thập phân.",
    "   B02: số âm trong ngoặc (1.234); các báo cáo khác: -1.234. Ngày: dd/mm/yyyy. Dùng hàm F_RPT_SO trong file SQL.",
    "6. Giới hạn engine: tối đa 200 biến / template; mỗi giá trị tối đa 50.000 ký tự (CĐTK cấp III ~800 dòng ≈ 35.000 ký tự / cột tên).",
    "   Nếu vượt -> tách theo đơn vị hoặc theo phần A/B.",
    "7. Cột 'Khóa lọc' không in ra, dùng để lọc / chọn dòng trong Report Builder.",
    "8. Bảng nguồn RPT_BMxx_HDR / RPT_BMxx_CT và view V_RPT_BMxx là ĐỀ XUẤT; team data có thể dựng khác, chỉ cần view cuối cùng trả đúng tên cột.",
]


def write_excel(results):
    wb = Workbook()
    ws = wb.active
    ws.title = "Huong_dan"
    for i, t in enumerate(GUIDE, 1):
        ws.cell(i, 1, t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 140

    ws = wb.create_sheet("Tong_quan")
    put_table(ws, 1, ["Mã", "Báo cáo", "File template", "Dataset (view)", "Hạt dữ liệu", "Loại", "Số biến template", "Bảng nguồn đề xuất"],
              [(r["code"], r["name"], r["file"], f"V_RPT_{r['code']}", r["grain"], "Cố định (pivot)" if r["kind"] == "fixed" else "Danh sách (nhiều dòng)",
                results[r["code"]]["n_vars"], f"RPT_{r['code']}_HDR, RPT_{r['code']}_CT") for r in REPORTS],
              [8, 42, 46, 16, 48, 22, 10, 36])

    for r in REPORTS:
        row = results[r["code"]]["row"]
        ws = wb.create_sheet(f"{r['code']}_Dataset")
        ws.cell(1, 1, f"{r['code']}. {r['name']} - Dataset V_RPT_{r['code']} ({r['grain']})").font = Font(bold=True, size=13)
        put_table(ws, 3, ["STT", "Cột dataset = Biến template", "Kiểu", "Mô tả", "Quy tắc", "Giá trị mẫu"],
                  [(i, n, t, d, rule, preview(row.get(n))) for i, (n, t, d, rule) in enumerate(dataset_columns(r), 1)],
                  [6, 22, 14, 60, 40, 60])
        ws.freeze_panes = "C4"

        ws = wb.create_sheet(f"{r['code']}_Nguon")
        ws.cell(1, 1, f"{r['code']} - Bảng nguồn đề xuất (team data đổ dữ liệu vào đây; view V_RPT_{r['code']} chuyển thành dataset)").font = Font(bold=True, size=13)
        put_table(ws, 3, ["Bảng", "Cột", "Kiểu", "Khóa", "Mô tả"], source_tables(r), [22, 20, 16, 8, 70])

        ws = wb.create_sheet(f"{r['code']}_Mau")
        ws.cell(1, 1, f"{r['code']} - Dữ liệu mẫu bảng chi tiết RPT_{r['code']}_CT (trích từ báo cáo gốc)").font = Font(bold=True, size=13)
        detail = results[r["code"]]["detail"]
        if r["kind"] == "fixed":
            hdrs = ["ma_chi_tieu", "stt", "ten_chi_tieu"] + [src for _, src, _ in r["measures"]]
            rows = [(f"{r['prefix']}_r{i + 1:02d}", stt, label, *vals) for i, ((stt, label, _), vals) in enumerate(zip(r["rows"], detail))]
        elif r["code"] == "BM04":
            hdrs = ["phan", "stt", "ten_tk", "tk_nhnn", "tk_fcc"] + AMT6
            rows = [(p.upper(), i, *[x[k] for k in hdrs[2:]]) for p in ("a", "b") for i, x in enumerate(detail[p], 1)]
        else:
            hdrs = list(detail[0].keys())
            rows = [tuple(x.values()) for x in detail]
        put_table(ws, 3, hdrs, rows, [16] * len(hdrs))
        for line in ws.iter_rows(min_row=4):
            for c in line:
                if isinstance(c.value, (int, float)):
                    c.number_format = "#,##0.00" if isinstance(c.value, float) else "#,##0"
    wb.save(OUT / "Cau_truc_du_lieu.xlsx")


# ================================================================ điền thử bằng engine
def render():
    SAMPLES.mkdir(exist_ok=True)
    results = {}
    for r in REPORTS:
        row, detail = sample_row(r)
        tpl = rt_office.OfficeTemplate((OUT / r["file"]).read_bytes(), "docx")
        cols = [n for n, *_ in dataset_columns(r)]
        missing = set(tpl.variables) - set(cols)
        unused = set(cols) - set(tpl.variables) - {k for k, *_ in r["keys"]}
        assert not missing and not unused, (r["code"], missing, unused)
        assert set(cols) == set(row), (r["code"], set(cols) ^ set(row))
        longest = max(len(str(v)) for v in row.values())
        assert longest <= 50000, (r["code"], longest)
        out = SAMPLES / f"{r['code']}_sample.docx"
        out.write_bytes(tpl.merge(row))
        (SAMPLES / f"{r['code']}.json").write_text(json.dumps(row, ensure_ascii=False, indent=1), encoding="utf-8")
        results[r["code"]] = {"row": row, "detail": detail, "n_vars": len(tpl.variables)}
        print(f"{r['code']}: {len(tpl.variables)} biến, giá trị dài nhất {longest} ký tự")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(SAMPLES), *map(str, sorted(SAMPLES.glob("*.docx")))],
                   check=True, capture_output=True)
    return results


if __name__ == "__main__":
    res = render()
    write_sql()
    write_excel(res)
    print("Xong.")
